"""
Export all pipeline outputs into a single, decision-ready Excel workbook
with styled headers, conditional risk-color highlighting, frozen panes and
auto-filters — the kind of deliverable a planning/ops team can act on
directly without touching a line of code.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
BODY_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(bold=True, color="1F4E78", name="Calibri", size=14)
SECTION_FONT = Font(bold=True, color="1F4E78", name="Calibri", size=11)
ITALIC_FONT = Font(italic=True, color="666666", name="Calibri", size=10)

RISK_FILLS = {
    "Overflow": "FF9999", "High risk": "FFD699", "Healthy": "C6E0B4",
    "Underutilized": "DDEBF7", "Surplus": "F2F2F2",
    "Closed (demand deferred to next open day)": "D9D9D9", "Closed": "E7E6E6",
}


def write_df(ws, df: pd.DataFrame, pct_cols=None, risk_col=None):
    pct_cols = pct_cols or []
    ws.append(list(df.columns))
    header_row = ws.max_row
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for _, row in df.iterrows():
        ws.append(list(row))

    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            if df.columns[c - 1] in pct_cols:
                cell.number_format = "0.0%"
        if risk_col and risk_col in df.columns:
            val = str(ws.cell(row=r, column=list(df.columns).index(risk_col) + 1).value)
            for key, color in RISK_FILLS.items():
                if key in val:
                    fill = PatternFill("solid", start_color=color)
                    for c in range(1, len(df.columns) + 1):
                        ws.cell(row=r, column=c).fill = fill
                    break

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    ws.auto_filter.ref = f"{get_column_letter(1)}{header_row}:{get_column_letter(len(df.columns))}{ws.max_row}"
    for c in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(15, len(str(df.columns[c - 1])) + 2)


def write_text_block(ws, lines):
    """lines: list of (text, font) tuples."""
    wrap = Alignment(wrap_text=True, vertical="top")
    for text, font in lines:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = font
        ws.cell(row=ws.max_row, column=1).alignment = wrap


def build_workbook(*, forecast_output, network_daily, dc_monthly, at_risk, surplus,
                    risk_by_type, backtest_summary, undelivered, scenario,
                    on_time_by_dc, network_on_time_rate, realloc_plan, backtest_daily,
                    multi_scenario, region_stats, methodology_lines,
                    management_summary_lines, output_path):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Forecast Output"
    disp = forecast_output[["Delivery Date", "DC ID", "DC Name", "Time Scope",
                             "Forecast Delivered Items", "Forecast Delivered Orders",
                             "Est. Staff Needed", "Capacity", "Capacity Utilization",
                             "Risk Category"]].rename(columns={"Capacity Utilization": "Utilization"})
    write_df(ws1, disp, pct_cols=["Utilization"], risk_col="Risk Category")

    ws2 = wb.create_sheet("Network Daily Summary")
    write_df(ws2, network_daily, pct_cols=["Utilization"])

    ws3 = wb.create_sheet("DC Monthly Summary")
    write_df(ws3, dc_monthly.sort_values("avg_utilization", ascending=False)[
        ["DC ID", "DC Name", "Service Type", "State", "active_days", "total_staff_needed",
         "avg_utilization", "max_utilization", "month_utilization", "days_high_risk",
         "days_overflow", "days_surplus", "pct_days_high_risk", "pct_days_surplus"]],
        pct_cols=["avg_utilization", "max_utilization", "month_utilization",
                  "pct_days_high_risk", "pct_days_surplus"])

    ws4 = wb.create_sheet("At-Risk DCs")
    write_df(ws4, at_risk[["DC ID", "DC Name", "Service Type", "State", "active_days",
                            "total_staff_needed", "avg_utilization", "max_utilization",
                            "days_overflow", "days_high_risk", "pct_days_high_risk"]],
             pct_cols=["avg_utilization", "max_utilization", "pct_days_high_risk"])

    ws5 = wb.create_sheet("Surplus Capacity DCs")
    write_df(ws5, surplus[["DC ID", "DC Name", "Service Type", "State", "active_days",
                            "avg_utilization", "month_utilization", "days_surplus", "pct_days_surplus"]],
             pct_cols=["avg_utilization", "month_utilization", "pct_days_surplus"])

    ws6 = wb.create_sheet("Risk by Service Type")
    write_df(ws6, risk_by_type, pct_cols=["Avg_Utilization", "pct_at_risk", "pct_surplus"])

    ws7 = wb.create_sheet("Validation - Backtest")
    ws7.append(["Leave-one-day-out cross-validation, WITH capacity capping applied to predictions."])
    ws7.cell(row=1, column=1).font = ITALIC_FONT
    ws7.append([])
    write_df(ws7, backtest_summary)
    ws7.append([])
    ws7.append(["Metric interpretation"])
    ws7.cell(row=ws7.max_row, column=1).font = SECTION_FONT
    for t in [
        "Bias < 0: model tends to under-forecast",
        "P90 AE: error threshold covering 90% of observations",
        "Hit Rate ±20%: share of forecasts within ±20% of actuals",
        "Under Forecast %: share of slots where the model under-predicted (operational shortage risk)",
        "Demand Accuracy: raw forecasting quality, before capacity constraints",
        "Capacity Accuracy: performance after operational constraints — reflects what actually ships",
    ]:
        ws7.append([t])

    ws8 = wb.create_sheet("Undelivered at Horizon End")
    ws8.append(["DC / time-window slots with deferred demand still unresolved at the end of the forecast horizon."])
    ws8.cell(row=1, column=1).font = ITALIC_FONT
    ws8.append([])
    if len(undelivered):
        write_df(ws8, undelivered)
    else:
        ws8.append(["(none)"])

    ws9 = wb.create_sheet("What-If Scenario Analysis")
    ws9.column_dimensions["A"].width = 30
    ws9.column_dimensions["B"].width = 55
    target = scenario["summary"].iloc[0, 1]
    ws9.append([f"WHAT-IF SCENARIO: Capacity Expansion for {target}"])
    ws9.cell(row=1, column=1).font = TITLE_FONT
    ws9.append([scenario["summary"].iloc[1, 1]])
    ws9.cell(row=2, column=1).font = ITALIC_FONT
    ws9.append([])
    ws9.append(["NETWORK IMPACT SUMMARY"])
    ws9.cell(row=ws9.max_row, column=1).font = SECTION_FONT
    for _, row in scenario["summary"].iterrows():
        ws9.append(list(row))
        for c in range(1, 3):
            ws9.cell(row=ws9.max_row, column=c).font = BODY_FONT
    ws9.append([])
    ws9.append(["COST / BENEFIT"])
    ws9.cell(row=ws9.max_row, column=1).font = SECTION_FONT
    for _, row in scenario["cost_benefit"].iterrows():
        ws9.append(list(row))
    ws9.append([])
    ws9.append(["DAILY BREAKDOWN"])
    ws9.cell(row=ws9.max_row, column=1).font = SECTION_FONT
    ws9.append([])
    detail_cols = [c for c in scenario["detail"].columns]
    write_df(ws9, scenario["detail"][detail_cols])

    ws10 = wb.create_sheet("On-Time Delivery Rate")
    ws10.append([f"Network-wide on-time delivery rate (same-day, not deferred from a closed day): {network_on_time_rate*100:.1f}%"])
    ws10.cell(row=1, column=1).font = SECTION_FONT
    ws10.append(["A new metric surfaced from data the original model already computed per row (same-day vs. "
                 "deferred demand) but did not previously roll up into its own service-level KPI."])
    ws10.cell(row=2, column=1).font = ITALIC_FONT
    ws10.append([])
    write_df(ws10, on_time_by_dc, pct_cols=["on_time_rate"])

    ws11 = wb.create_sheet("Staff Reallocation Plan")
    ws11.append(["Greedy pairing of each at-risk DC with the best-fitting surplus DC (same Service Type "
                 "preferred), sized to close roughly half the staffing gap -- a conservative first move."])
    ws11.cell(row=1, column=1).font = ITALIC_FONT
    ws11.append([])
    if len(realloc_plan):
        write_df(ws11, realloc_plan, pct_cols=["At-Risk Utilization", "Donor Utilization"])
    else:
        ws11.append(["(no at-risk/surplus DC pairs to reallocate between)"])

    ws12 = wb.create_sheet("Backtest Daily Trend")
    ws12.append(["Day-by-day network-wide actual vs. capacity-constrained predicted volume -- "
                 "the time series the aggregate backtest metrics on 'Validation - Backtest' were computed from."])
    ws12.cell(row=1, column=1).font = ITALIC_FONT
    ws12.append([])
    write_df(ws12, backtest_daily)

    ws13 = wb.create_sheet("Multi-DC What-If Comparison")
    ws13.append(["The single-DC What-If scenario, re-run for the top at-risk DCs and ranked by net economic "
                 "benefit -- turns one scenario into a prioritized capital-investment shortlist."])
    ws13.cell(row=1, column=1).font = ITALIC_FONT
    ws13.append([])
    if len(multi_scenario):
        write_df(ws13, multi_scenario, pct_cols=["Pre-Scenario Avg. Utilization"])
    else:
        ws13.append(["(no at-risk DCs to compare)"])

    ws14 = wb.create_sheet("Region Summary")
    write_df(ws14, region_stats, pct_cols=["Avg_Utilization"])

    ws_meth = wb.create_sheet("Methodology")
    ws_meth.column_dimensions["A"].width = 120
    write_text_block(ws_meth, methodology_lines)

    ws_mgmt = wb.create_sheet("Management Summary")
    ws_mgmt.column_dimensions["A"].width = 120
    write_text_block(ws_mgmt, management_summary_lines)

    order = [
        "Methodology", "Management Summary", "Forecast Output", "Network Daily Summary",
        "DC Monthly Summary", "At-Risk DCs", "Surplus Capacity DCs", "Risk by Service Type",
        "Region Summary", "On-Time Delivery Rate", "Staff Reallocation Plan",
        "Validation - Backtest", "Backtest Daily Trend", "Undelivered at Horizon End",
        "What-If Scenario Analysis", "Multi-DC What-If Comparison",
    ]
    wb._sheets = [wb[name] for name in order]

    wb.save(output_path)
    return output_path
